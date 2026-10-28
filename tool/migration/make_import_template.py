"""Builds migration/Entre Sales - data import template.xlsx: the workbook the
owners fill in with the store's real artists, products/stock and past sales.

    python tool/migration/make_import_template.py
"""
import os, sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, os.path.join(ROOT, 'tool', 'demo_data'))
from catalog_mapping import CATEGORIES  # same list the database starts with

OUT = os.path.join(ROOT, 'migration', 'Entre Sales - data import template.xlsx')
ROWS = 3000  # rows prepared with dropdowns / formats

FONT = 'Arial'
F = lambda **k: Font(name=FONT, size=10, **k)
REQ_FILL = PatternFill('solid', fgColor='5B3A2C')      # required column header
OPT_FILL = PatternFill('solid', fgColor='B08B78')      # optional column header
CALC_FILL = PatternFill('solid', fgColor='9E9E9E')     # filled in by formula
INPUT_FILL = PatternFill('solid', fgColor='FFF8DC')    # cells to fill in
EXAMPLE_FILL = PatternFill('solid', fgColor='E7E6E6')  # the example row
THIN = Side(style='thin', color='D0D0D0')
GENDERS = ['Male', 'Female', 'Queer']
FIELDS = ['Size', 'Style', 'Design', 'Finish', 'Pattern', 'Scent', 'Language', 'Option']
assert {f for _, _, a in CATEGORIES for f in a} <= set(FIELDS)

wb = Workbook()


def table(ws, columns, example, widths):
    """columns: (header, required, note). Row 1 headers, row 2 example, data from row 3."""
    for i, (head, required, note) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=head + (' *' if required == True else ''))
        c.font = F(bold=True, color='FFFFFF')
        c.fill = CALC_FILL if required == 'calc' else (REQ_FILL if required else OPT_FILL)
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.comment = Comment(note, 'Entre Sales')
        c.comment.width, c.comment.height = 320, 140
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[1].height = 42
    for i, v in enumerate(example, start=1):
        c = ws.cell(row=2, column=i, value=v)
        c.fill = EXAMPLE_FILL
        c.font = F(italic=True, color='595959')
    for r in range(3, ROWS + 3):
        for i, (_, required, _) in enumerate(columns, start=1):
            c = ws.cell(row=r, column=i)
            c.font = F()
            if required != 'calc':
                c.fill = INPUT_FILL
            c.border = Border(bottom=THIN)
    ws.freeze_panes = 'B3' if ws.title != 'Artists' else 'B3'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{ROWS + 2}'


def dropdown(ws, col, source, prompt, error, allow_blank=True):
    dv = DataValidation(type='list', formula1=source, allow_blank=allow_blank, showErrorMessage=True,
                        errorTitle='Not in the list', error=error, promptTitle='Choose', prompt=prompt)
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}{ROWS + 2}')


def number(ws, col, kind, lo, hi, error):
    dv = DataValidation(type=kind, operator='between', formula1=str(lo), formula2=str(hi),
                        allow_blank=True, showErrorMessage=True, errorTitle='Check the number', error=error)
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}{ROWS + 2}')


def fmt(ws, col, number_format):
    for r in range(2, ROWS + 3):
        ws[f'{col}{r}'].number_format = number_format


# ------------------------------------------------------------- Start here
ws = wb.active
ws.title = 'Start here'
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 26
ws.column_dimensions['C'].width = 95
ws.sheet_properties.pageSetUpPr.fitToPage = True  # prints one page wide
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
lines = [
    ('title', 'Entre Sales — data to bring into the new system'),
    ('text', 'Fill in the three yellow tabs. Everything is imported in one go on the switch-over day.'),
    ('gap', None),
    ('head', 'The tabs'),
    ('row', ('Artists', 'One row per artist, with billing details and commission.')),
    ('row', ('Products', 'One row per variant (each size / colour / design) with its price and the stock on the shelf.')),
    ('row', ('Past sales', 'One row per item on each receipt, from the old system, up to the switch-over.')),
    ('row', ('Categories', 'Reference only: the categories, their VAT and which fields (Size, Design…) each uses.')),
    ('gap', None),
    ('head', 'How to read the colours'),
    ('legend', ('Dark header  *', 'Required for every row.', REQ_FILL)),
    ('legend', ('Light header', 'Optional: leave empty if you don\'t have it.', OPT_FILL)),
    ('legend', ('Grey header', 'Calculated for you (a check). Don\'t type here.', CALC_FILL)),
    ('legend', ('Yellow cells', 'Where you type.', INPUT_FILL)),
    ('legend', ('Grey row 2', 'An example of a filled-in row. It is skipped by the import; you can leave it.', EXAMPLE_FILL)),
    ('row', ('Header notes', 'Hover over (or tap) any column header for what it means.')),
    ('gap', None),
    ('head', 'Formats'),
    ('row', ('Money', 'Euros with decimals, e.g. 25 or 25.50. No € sign needed.')),
    ('row', ('Percentages', 'Plain numbers: 35 means 35%. Leave empty to use the default.')),
    ('row', ('Dates / times', 'Day/month/year, e.g. 12/09/2026, and 24-hour time, e.g. 17:42.')),
    ('row', ('Yes / No and lists', 'Pick from the dropdown in the cell so the spelling matches.')),
    ('row', ('Names', 'Write each artist and product name exactly the same way on every tab.')),
    ('gap', None),
    ('head', 'Switch-over day'),
    ('row', ('1. Close the old till', 'Stop selling in the old system (e.g. after closing).')),
    ('row', ('2. Count the stock', 'Enter what is physically on the shelf in Products → Stock now. Past sales do not lower it again.')),
    ('row', ('3. Last sales', 'Make sure Past sales has every receipt up to the close.')),
    ('row', ('4. Import', 'The demo data is cleared and this workbook is imported. From then on, sell only in the new system.')),
    ('gap', None),
    ('head', 'Good to know'),
    ('row', ('Refunded sales', 'Past refunds are not imported. Leave fully refunded sales out, and for part-refunds enter only what was kept.')),
    ('row', ('Sold-out products', 'Past sales can include products no longer in stock; they don\'t need a row in Products.')),
    ('row', ('Store details', 'Store name, address and NIF are entered in the app: Settings → Store details.')),
    ('row', ('Commission on old sales', 'Uses the rate in the Commission column of Past sales if filled, otherwise the artist\'s (or product\'s) rate.')),
]
r = 1
for kind, v in lines:
    r += 1
    if kind == 'title':
        ws.cell(row=r, column=2, value=v).font = Font(name=FONT, size=16, bold=True, color='5B3A2C')
    elif kind == 'text':
        ws.cell(row=r, column=2, value=v).font = F()
    elif kind == 'head':
        ws.cell(row=r, column=2, value=v).font = Font(name=FONT, size=12, bold=True, color='5B3A2C')
    elif kind == 'row':
        ws.cell(row=r, column=2, value=v[0]).font = F(bold=True)
        c = ws.cell(row=r, column=3, value=v[1]); c.font = F(); c.alignment = Alignment(wrap_text=True)
    elif kind == 'legend':
        c = ws.cell(row=r, column=2, value=v[0]); c.fill = v[2]
        c.font = F(bold=True, color='FFFFFF' if v[2] in (REQ_FILL, OPT_FILL, CALC_FILL) else '000000')
        ws.cell(row=r, column=3, value=v[1]).font = F()

# ------------------------------------------------------------------ Lists
lists = wb.create_sheet('Lists')
for col, (head, values) in enumerate([
    ('Yes/No', ['Yes', 'No']),
    ('Gender', GENDERS),
    ('Payment', ['Cash', 'Card']),
    ('Category', [c[0] for c in CATEGORIES]),
    ('Discount', ['Friends & Family']),
], start=1):
    lists.cell(row=1, column=col, value=head).font = F(bold=True)
    for i, v in enumerate(values, start=2):
        lists.cell(row=i, column=col, value=v).font = F()
lists.sheet_state = 'hidden'
LIST = lambda col, n: f'=Lists!${col}$2:${col}${n + 1}'

# ---------------------------------------------------------------- Artists
ws = wb.create_sheet('Artists')
ws.sheet_properties.tabColor = 'F4C430'
table(ws, [
    ('Artist name', True, 'The name used in the store and on products. Must be unique, and written the same way on the other tabs.'),
    ('Spain resident', True, 'Yes = lives in Spain (the 10% resident tax is added on top of their payout). No = lives abroad.'),
    ('Commission (%)', False, 'The store\'s commission for this artist, e.g. 35. Leave empty for the default 35%.'),
    ('Gender', False, 'For sales analysis: Male, Female or Queer.'),
    ('Legal name', False, 'Name for invoices / bank transfers, if different from the artist name.'),
    ('NIF / NIE / Tax ID', False, 'Tax ID printed on their payout statement.'),
    ('Address', False, 'Billing address.'),
    ('IBAN', False, 'Bank account for payouts, e.g. ES12 3456 7890 1234 5678 9012.'),
    ('Bank', False, 'Bank name.'),
], ['Jumbo Press', 'Yes', None, 'Female', 'Jumbo Press S.L.', 'B12345678', 'Calle Ejemplo 1, 28001 Madrid',
    'ES12 3456 7890 1234 5678 9012', 'Banco Ejemplo'],
    [26, 15, 15, 12, 26, 18, 36, 32, 18])
dropdown(ws, 'B', LIST('A', 2), 'Yes or No', 'Choose Yes or No.', allow_blank=True)
dropdown(ws, 'D', LIST('B', len(GENDERS)), 'Male, Female or Queer', 'Choose from the list.')
number(ws, 'C', 'decimal', 0, 100, 'A percentage from 0 to 100, e.g. 35.')
ws.conditional_formatting.add(f'A3:A{ROWS + 2}', FormulaRule(
    formula=['AND(A3<>"",COUNTIF($A$3:$A$' + str(ROWS + 2) + ',A3)>1)'], fill=PatternFill('solid', fgColor='F4B6B6')))

# --------------------------------------------------------------- Products
ws = wb.create_sheet('Products')
ws.sheet_properties.tabColor = 'F4C430'
cols = [
    ('Artist', True, 'Pick from the artists on the Artists tab.'),
    ('Product', True, 'Product name. Rows with the same Artist + Product are one product with several variants.'),
    ('Category', True, 'Decides the VAT and which fields apply (see the Categories tab).'),
]
for f in FIELDS:
    users = ', '.join(c[0] for c in CATEGORIES if f in c[2])
    cols.append((f, False, f'Fill only if this product\'s category uses {f} ({users}). One value per row; '
                           f'each different combination of fields is its own row (variant).'))
cols += [
    ('Price incl. VAT (€)', True, 'Shelf price the customer pays, VAT included. The price before VAT is calculated from it.'),
    ('Stock now', True, 'Units on the shelf on the switch-over day (0 if sold out but still offered).'),
    ('Date arrived', False, 'When this stock arrived (for "days to sell" reports). Empty = the switch-over day.'),
    ('Product commission (%)', False, 'Only if this product has a different commission from its artist, e.g. 30. Empty = the artist\'s rate.'),
    ('VAT (%) override', False, 'Only for a product with a different VAT from its category. Empty = the category\'s VAT.'),
]
example = ['Jumbo Press', 'Risograph Poster', 'Prints & Posters', 'A3', None, 'Bathroom', None, None, None, None, None,
           25, 4, None, None, None]
table(ws, cols, example, [22, 28, 20] + [11] * len(FIELDS) + [14, 11, 13, 14, 13])
fc = len(FIELDS)
L = lambda i: get_column_letter(i)
price, stock, arrived, pcomm, vat = (L(4 + fc + i) for i in range(5))
dropdown(ws, 'A', f'=Artists!$A$3:$A${ROWS + 2}', 'Artists from the Artists tab', 'Add the artist on the Artists tab first.')
dropdown(ws, 'C', LIST('D', len(CATEGORIES)), 'Choose the category', 'Choose a category from the list.', allow_blank=False)
number(ws, price, 'decimal', 0, 100000, 'A price in euros, e.g. 25 or 25.50.')
number(ws, stock, 'whole', 0, 100000, 'A whole number of units, 0 or more.')
number(ws, pcomm, 'decimal', 0, 100, 'A percentage from 0 to 100.')
number(ws, vat, 'decimal', 0, 100, 'A percentage from 0 to 100.')
dv = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True, showErrorMessage=True,
                    errorTitle='Check the date', error='A date, e.g. 12/09/2026.')
ws.add_data_validation(dv); dv.add(f'{arrived}2:{arrived}{ROWS + 2}')
fmt(ws, price, '#,##0.00 "€"'); fmt(ws, arrived, 'DD/MM/YYYY')
# Grey out field cells the chosen category doesn't use.
for i, f in enumerate(FIELDS):
    col = L(4 + i)
    using = [c[0] for c in CATEGORIES if f in c[2]]
    cond = 'AND($C3<>"",' + ('NOT(OR(' + ','.join(f'$C3="{u}"' for u in using) + '))' if using else 'TRUE') + ')'
    ws.conditional_formatting.add(f'{col}3:{col}{ROWS + 2}',
                                  FormulaRule(formula=[cond], fill=PatternFill('solid', fgColor='D9D9D9')))

# ------------------------------------------------------------- Past sales
ws = wb.create_sheet('Past sales')
ws.sheet_properties.tabColor = 'F4C430'
cols = [
    ('Receipt no.', True, 'The receipt / ticket number in the old system. All items of one receipt share it.'),
    ('Date', True, 'Day of the sale, e.g. 12/09/2026.'),
    ('Time', False, 'Time of the sale, e.g. 17:42 (for the busiest-hours report). Empty = 12:00.'),
    ('Payment', True, 'Cash or Card.'),
    ('Artist', True, 'Whose product it was (the Artists tab name).'),
    ('Product', True, 'Product name. If it is on the Products tab, write it the same way.'),
    ('Category', True, 'The product\'s category (for reports by category).'),
    ('Variant', False, 'Size / colour / design as on the receipt, e.g. "A3 · Bathroom". Empty if the product has no variants.'),
    ('Quantity', True, 'Units of this item on the receipt.'),
    ('Total before VAT (€)', True, 'What this line came to before VAT (quantity × price, after any discount).'),
    ('Total incl. VAT (€)', True, 'What the customer paid for this line, VAT included.'),
    ('VAT % (check)', 'calc', 'Calculated from the two totals. Should read 21%, 10% or 4%; anything else means a typo.'),
    ('Discount', False, 'Friends & Family if it applied to this receipt; otherwise empty.'),
    ('Commission (%)', False, 'Only if this sale had a different commission from the artist\'s / product\'s rate today.'),
]
table(ws, cols, ['T-1024', None, None, 'Card', 'Jumbo Press', 'Risograph Poster', 'Prints & Posters', 'A3 · Bathroom',
                 2, 41.32, 50, None, None, None], [12, 12, 8, 9, 22, 26, 20, 20, 9, 14, 14, 11, 16, 14])
from datetime import date, time
ws['B2'] = date(2026, 9, 12); ws['C2'] = time(17, 42)
dropdown(ws, 'D', LIST('C', 2), 'Cash or Card', 'Choose Cash or Card.', allow_blank=False)
dropdown(ws, 'E', f'=Artists!$A$3:$A${ROWS + 2}', 'Artists from the Artists tab', 'Add the artist on the Artists tab first.')
dropdown(ws, 'G', LIST('D', len(CATEGORIES)), 'Choose the category', 'Choose a category from the list.')
dropdown(ws, 'M', LIST('E', 1), 'Friends & Family, or empty', 'Choose from the list or leave empty.')
number(ws, 'I', 'whole', 1, 10000, 'A whole number, 1 or more.')
number(ws, 'J', 'decimal', 0, 1000000, 'An amount in euros.')
number(ws, 'K', 'decimal', 0, 1000000, 'An amount in euros.')
number(ws, 'N', 'decimal', 0, 100, 'A percentage from 0 to 100.')
dv = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True, showErrorMessage=True,
                    errorTitle='Check the date', error='A date, e.g. 12/09/2026.')
ws.add_data_validation(dv); dv.add(f'B2:B{ROWS + 2}')
fmt(ws, 'B', 'DD/MM/YYYY'); fmt(ws, 'C', 'HH:MM'); fmt(ws, 'J', '#,##0.00 "€"'); fmt(ws, 'K', '#,##0.00 "€"')
for r in range(2, ROWS + 3):
    ws[f'L{r}'] = f'=IF(OR(J{r}="",K{r}="",J{r}=0),"",K{r}/J{r}-1)'
    ws[f'L{r}'].number_format = '0%'
    ws[f'L{r}'].font = F(color='595959')
# Flag a VAT that isn't 0 / 4 / 10 / 21 % (allowing for cent rounding).
ws.conditional_formatting.add(f'L3:L{ROWS + 2}', FormulaRule(
    formula=['AND(L3<>"",ABS(L3-0)>0.005,ABS(L3-0.04)>0.005,ABS(L3-0.1)>0.005,ABS(L3-0.21)>0.005)'],
    fill=PatternFill('solid', fgColor='F4B6B6')))

# ------------------------------------------------------------- Categories
ws = wb.create_sheet('Categories')
ws.sheet_properties.tabColor = '9E9E9E'
heads = ['Category', 'VAT (%)', 'Fields used', 'Fixed options']
for i, h in enumerate(heads, start=1):
    c = ws.cell(row=1, column=i, value=h); c.font = F(bold=True, color='FFFFFF'); c.fill = CALC_FILL
for r, (name, tax, attrs) in enumerate(CATEGORIES, start=2):
    ws.cell(row=r, column=1, value=name).font = F()
    ws.cell(row=r, column=2, value=tax).font = F()
    ws.cell(row=r, column=3, value=', '.join(attrs) or '—').font = F()
    ws.cell(row=r, column=4, value='; '.join(f'{k}: {", ".join(v)}' for k, v in attrs.items() if v) or '—').font = F()
note = ws.cell(row=len(CATEGORIES) + 3, column=1,
               value='These are the categories in the system now. To change one (name, VAT, fields), tell the developer '
                     'before filling in Products, or change it in the app under Settings → Categories.')
note.font = F(italic=True, color='595959')
for col, w in zip('ABCD', [24, 9, 22, 50]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'

wb.move_sheet('Lists', offset=10)  # keep the hidden list last
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('written', os.path.relpath(OUT, ROOT))
